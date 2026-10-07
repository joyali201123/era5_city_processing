#!/usr/bin/env python3
"""按 10°×10° 格网下载 ERA5-Land 小时级 2 m 温度并维护 Excel 台账。"""

from __future__ import annotations

import argparse
import calendar
import logging
import os
import shutil
import sys
import threading
import time
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Iterable

import cdsapi
import xarray as xr
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


DATASET = "reanalysis-era5-land"
VARIABLE = "2m_temperature"
GRID_DEGREES = 10
DEFAULT_START_YEAR = 1950
DEFAULT_END_YEAR = 2026
DEFAULT_AVAILABILITY_LAG_DAYS = 5
DEFAULT_WORKERS = 2
DEFAULT_RETRIES = 3
DEFAULT_FLUSH_SECONDS = 60

TASK_HEADERS = [
    "task_id", "year", "grid_id", "row", "col",
    "north", "south", "west", "east",
    "request_north", "request_south", "request_west", "request_east",
    "status", "attempts", "start_time", "end_time", "duration_seconds",
    "file_size_bytes", "coverage_start", "coverage_end", "output_file",
    "last_error", "updated_at",
]

STATUS_PENDING = "pending"
STATUS_DOING = "doing"
STATUS_COMPLETED = "completed"
STATUS_PARTIAL = "partial"
STATUS_FAILED = "failed"


@dataclass(frozen=True)
class Grid:
    grid_id: str
    row: int
    col: int
    north: int
    south: int
    west: int
    east: int
    request_north: float
    request_south: float
    request_west: float
    request_east: float


@dataclass(frozen=True)
class Task:
    year: int
    grid: Grid
    output_file: Path

    @property
    def task_id(self) -> str:
        return f"{self.year}_{self.grid.grid_id}"


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def iso_utc(value: datetime | None = None) -> str:
    return (value or utc_now()).replace(microsecond=0).isoformat()


def coord_label(value: int, positive: str, negative: str, width: int) -> str:
    prefix = positive if value >= 0 else negative
    return f"{prefix}{abs(value):0{width}d}"


def tile_filename(year: int, grid: Grid) -> str:
    return (
        f"ERA5Land_t2m_{year}_{grid.grid_id}_"
        f"{coord_label(grid.north, 'N', 'S', 2)}_"
        f"{coord_label(grid.south, 'N', 'S', 2)}_"
        f"{coord_label(grid.west, 'E', 'W', 3)}_"
        f"{coord_label(grid.east, 'E', 'W', 3)}.nc"
    )


def build_grids() -> list[Grid]:
    """生成 648 个格子；请求边界避免相邻格重复 0.1° 网格点。"""
    grids: list[Grid] = []
    number = 1
    for row in range(18):
        north = 90 - row * GRID_DEGREES
        south = north - GRID_DEGREES
        request_south = -90.0 if row == 17 else south + 0.1
        for col in range(36):
            west = -180 + col * GRID_DEGREES
            east = west + GRID_DEGREES
            grids.append(
                Grid(
                    grid_id=f"G{number:04d}",
                    row=row + 1,
                    col=col + 1,
                    north=north,
                    south=south,
                    west=west,
                    east=east,
                    request_north=float(north),
                    request_south=float(request_south),
                    request_west=float(west),
                    request_east=float(east - 0.1),
                )
            )
            number += 1
    return grids


def available_end_for_year(year: int, lag_days: int) -> date | None:
    cutoff = (utc_now() - timedelta(days=lag_days)).date()
    if year > cutoff.year:
        return None
    if year < cutoff.year:
        return date(year, 12, 31)
    return cutoff


def expected_hours(year: int, end_date: date) -> int:
    hours = ((end_date - date(year, 1, 1)).days + 1) * 24

    # ERA5-Land从1950-01-01 01:00开始，
    # 1950-01-01 00:00不存在。
    if year == 1950:
        hours -= 1

    return hours


def build_tasks(start_year: int, end_year: int, output_dir: Path) -> list[Task]:
    grids = build_grids()
    return [
        Task(year, grid, output_dir / str(year) / tile_filename(year, grid))
        for year in range(start_year, end_year + 1)
        for grid in grids
    ]


def make_request(year: int, grid: Grid, months: Iterable[int], days: Iterable[int]) -> dict:
    return {
        "variable": [VARIABLE],
        "year": str(year),
        "month": [f"{month:02d}" for month in months],
        "day": [f"{day:02d}" for day in days],
        "time": [f"{hour:02d}:00" for hour in range(24)],
        "data_format": "netcdf",
        "download_format": "unarchived",
        "area": [
            grid.request_north,
            grid.request_west,
            grid.request_south,
            grid.request_east,
        ],
    }


def request_chunks(
    year: int,
    end_date: date,
    grid: Grid,
) -> list[tuple[str, dict]]:
    """逐月请求，最后合并为年度NetCDF。"""
    chunks: list[tuple[str, dict]] = []

    for month in range(1, end_date.month + 1):
        last_day = calendar.monthrange(year, month)[1]

        if month == end_date.month:
            last_day = min(last_day, end_date.day)

        chunk_name = (
            f"month_{month:02d}_days_01_{last_day:02d}"
        )

        request = make_request(
            year=year,
            grid=grid,
            months=[month],
            days=range(1, last_day + 1),
        )

        chunks.append((chunk_name, request))

    return chunks

def find_time_name(ds: xr.Dataset) -> str:
    for name in ("valid_time", "time"):
        if name in ds.dims or name in ds.coords:
            return name
    raise ValueError("NetCDF 中找不到 valid_time 或 time 坐标")


def inspect_netcdf(path: Path) -> tuple[int, str, str]:
    with xr.open_dataset(path, decode_times=True) as ds:
        time_name = find_time_name(ds)
        count = int(ds.sizes.get(time_name, ds[time_name].size))
        if count <= 0:
            raise ValueError("NetCDF 时间维为空")
        first = str(ds[time_name].values[0])
        last = str(ds[time_name].values[-1])
        if "t2m" not in ds.data_vars and VARIABLE not in ds.data_vars:
            raise ValueError(f"NetCDF 缺少温度变量，现有变量：{list(ds.data_vars)}")
        return count, first, last


def validate_netcdf(path: Path, minimum_hours: int | None = None) -> tuple[int, str, str]:
    if not path.exists() or path.stat().st_size < 1024:
        raise ValueError("目标文件不存在或文件过小")
    count, first, last = inspect_netcdf(path)
    if minimum_hours is not None and count < minimum_hours:
        raise ValueError(f"小时数不足：得到 {count}，预期至少 {minimum_hours}")
    return count, first, last


def normalize_download(downloaded: Path) -> None:
    if not zipfile.is_zipfile(downloaded):
        return
    extract_dir = downloaded.with_suffix(downloaded.suffix + ".unzipped")
    if extract_dir.exists():
        shutil.rmtree(extract_dir)
    extract_dir.mkdir(parents=True)
    with zipfile.ZipFile(downloaded) as archive:
        candidates = [name for name in archive.namelist() if name.lower().endswith(".nc")]
        if len(candidates) != 1:
            raise ValueError(f"ZIP 中预期 1 个 NetCDF，实际为 {len(candidates)} 个")
        archive.extract(candidates[0], extract_dir)
    source = extract_dir / candidates[0]
    replacement = downloaded.with_suffix(downloaded.suffix + ".nc")
    shutil.move(str(source), replacement)
    downloaded.unlink(missing_ok=True)
    os.replace(replacement, downloaded)
    shutil.rmtree(extract_dir, ignore_errors=True)


def retrieve_to_file(client: cdsapi.Client, request: dict, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_suffix(target.suffix + ".downloading")
    temp.unlink(missing_ok=True)
    client.retrieve(DATASET, request, str(temp))
    normalize_download(temp)
    validate_netcdf(temp)
    os.replace(temp, target)


def merge_netcdf(parts: list[Path], target: Path) -> None:
    datasets: list[xr.Dataset] = []
    temp = target.with_suffix(target.suffix + ".merging")
    temp.unlink(missing_ok=True)
    try:
        for part in parts:
            datasets.append(xr.open_dataset(part, decode_times=True))
        time_name = find_time_name(datasets[0])
        combined = xr.concat(
            datasets,
            dim=time_name,
            data_vars="minimal",
            coords="minimal",
            compat="override",
        ).sortby(time_name)
        index = combined.indexes[time_name]
        if index.has_duplicates:
            combined = combined.isel({time_name: ~index.duplicated()})
        encoding = {
            name: {"zlib": True, "complevel": 4}
            for name in combined.data_vars
            if combined[name].dtype.kind in "fiu"
        }
        combined.to_netcdf(temp, engine="netcdf4", encoding=encoding)
        combined.close()
        os.replace(temp, target)
    finally:
        for ds in datasets:
            ds.close()
        temp.unlink(missing_ok=True)


def create_tracker_workbook(path: Path, tasks: list[Task], lag_days: int) -> None:
    wb = Workbook()
    overview = wb.active
    overview.title = "下载概览"
    task_sheet = wb.create_sheet("下载任务")
    grid_sheet = wb.create_sheet("格网索引")

    overview.append(["ERA5-Land 2米温度下载概览", ""])
    overview.append(["更新时间", iso_utc()])
    overview.append(["状态", "任务数"])
    overview.append(["全部任务", len(tasks)])
    overview.append([STATUS_PENDING, len(tasks)])
    overview.append([STATUS_DOING, 0])
    overview.append([STATUS_COMPLETED, 0])
    overview.append([STATUS_PARTIAL, 0])
    overview.append([STATUS_FAILED, 0])
    overview.append(["累计耗时（小时）", 0])
    overview.append(["平均耗时（分钟/任务）", 0])
    overview.append(["可用性延迟（天）", lag_days])
    overview.append(["状态说明", "pending=未下载，doing=下载中，completed=完整年度，partial=当年已下载至当前可用日期，failed=失败"])

    task_sheet.append(TASK_HEADERS)
    for task in tasks:
        end_date = available_end_for_year(task.year, lag_days)
        task_sheet.append([
            task.task_id, task.year, task.grid.grid_id, task.grid.row, task.grid.col,
            task.grid.north, task.grid.south, task.grid.west, task.grid.east,
            task.grid.request_north, task.grid.request_south,
            task.grid.request_west, task.grid.request_east,
            STATUS_PENDING, 0, "", "", "", "", date(task.year, 1, 1).isoformat(),
            end_date.isoformat() if end_date else "", str(task.output_file), "", iso_utc(),
        ])

    grid_headers = [
        "grid_id", "row", "col", "north", "south", "west", "east",
        "request_north", "request_south", "request_west", "request_east",
    ]
    grid_sheet.append(grid_headers)
    for grid in build_grids():
        grid_sheet.append([
            grid.grid_id, grid.row, grid.col, grid.north, grid.south, grid.west, grid.east,
            grid.request_north, grid.request_south, grid.request_west, grid.request_east,
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in (task_sheet, grid_sheet):
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.auto_filter.ref = sheet.dimensions

    overview.sheet_view.showGridLines = False
    overview["A1"].font = Font(size=14, bold=True, color="1F4E78")
    overview["A3"].fill = header_fill
    overview["B3"].fill = header_fill
    overview["A3"].font = header_font
    overview["B3"].font = header_font

    status_col = "N"
    last_row = task_sheet.max_row
    task_sheet[f"{status_col}2:{status_col}{last_row}"]
    colors = {
        STATUS_PENDING: ("FFF2CC", "7F6000"),
        STATUS_DOING: ("DDEBF7", "1F4E78"),
        STATUS_COMPLETED: ("E2F0D9", "375623"),
        STATUS_PARTIAL: ("FCE4D6", "843C0C"),
        STATUS_FAILED: ("F4CCCC", "9C0006"),
    }
    for value, (fill, font) in colors.items():
        task_sheet.conditional_formatting.add(
            f"{status_col}2:{status_col}{last_row}",
            FormulaRule(
                formula=[f'${status_col}2="{value}"'],
                fill=PatternFill("solid", fgColor=fill),
                font=Font(color=font, bold=True),
            ),
        )

    widths = {
        "A": 14, "B": 8, "C": 10, "D": 6, "E": 6, "F": 9, "G": 9,
        "H": 9, "I": 9, "J": 14, "K": 14, "L": 14, "M": 14, "N": 12,
        "O": 10, "P": 22, "Q": 22, "R": 18, "S": 18, "T": 14, "U": 14,
        "V": 65, "W": 60, "X": 22,
    }
    for col, width in widths.items():
        task_sheet.column_dimensions[col].width = width
    for col in "ABCDEFGHIJK":
        grid_sheet.column_dimensions[col].width = 15
    overview.column_dimensions["A"].width = 28
    overview.column_dimensions["B"].width = 105

    task_table = Table(displayName="DownloadTasks", ref=task_sheet.dimensions)
    task_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    task_sheet.add_table(task_table)
    grid_table = Table(displayName="GridIndex", ref=grid_sheet.dimensions)
    grid_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    grid_sheet.add_table(grid_table)

    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(".tmp.xlsx")
    wb.save(temp)
    os.replace(temp, path)


class Tracker:
    def __init__(self, path: Path, tasks: list[Task], lag_days: int, flush_seconds: int):
        self.path = path
        self.lock = threading.RLock()
        self.flush_seconds = flush_seconds
        self.last_save = 0.0
        if not path.exists():
            create_tracker_workbook(path, tasks, lag_days)
        self.wb = load_workbook(path)
        self.sheet = self.wb["下载任务"]
        self.overview = self.wb["下载概览"]
        headers = [cell.value for cell in self.sheet[1]]
        self.col = {name: index + 1 for index, name in enumerate(headers)}
        missing = [name for name in TASK_HEADERS if name not in self.col]
        if missing:
            raise ValueError(f"Excel 台账缺少字段：{missing}")
        self.rows: dict[str, int] = {}
        self.status_counts: Counter[str] = Counter()
        self.total_duration = 0.0
        self.duration_count = 0
        for row in range(2, self.sheet.max_row + 1):
            task_id = str(self.sheet.cell(row, self.col["task_id"]).value)
            self.rows[task_id] = row
            status = str(self.sheet.cell(row, self.col["status"]).value or STATUS_PENDING)
            self.status_counts[status] += 1
            duration = self.sheet.cell(row, self.col["duration_seconds"]).value
            if isinstance(duration, (int, float)):
                self.total_duration += float(duration)
                self.duration_count += 1

    def value(self, task_id: str, field: str):
        return self.sheet.cell(self.rows[task_id], self.col[field]).value

    def update(self, task_id: str, *, force_save: bool = False, **fields) -> None:
        with self.lock:
            row = self.rows[task_id]
            if "status" in fields:
                old_status = str(self.sheet.cell(row, self.col["status"]).value or STATUS_PENDING)
                new_status = str(fields["status"])
                if old_status != new_status:
                    self.status_counts[old_status] -= 1
                    self.status_counts[new_status] += 1
            if "duration_seconds" in fields:
                old_duration = self.sheet.cell(row, self.col["duration_seconds"]).value
                if isinstance(old_duration, (int, float)):
                    self.total_duration -= float(old_duration)
                    self.duration_count -= 1
                new_duration = fields["duration_seconds"]
                if isinstance(new_duration, (int, float)):
                    self.total_duration += float(new_duration)
                    self.duration_count += 1
            fields["updated_at"] = iso_utc()
            for field, value in fields.items():
                self.sheet.cell(row, self.col[field], value)
            self.maybe_save(force=force_save)

    def increment_attempts(self, task_id: str) -> int:
        with self.lock:
            attempts = int(self.value(task_id, "attempts") or 0) + 1
            self.update(task_id, attempts=attempts)
            return attempts

    def _write_summary(self) -> None:
        self.overview["B2"] = iso_utc()
        self.overview["B4"] = len(self.rows)
        self.overview["B5"] = self.status_counts[STATUS_PENDING]
        self.overview["B6"] = self.status_counts[STATUS_DOING]
        self.overview["B7"] = self.status_counts[STATUS_COMPLETED]
        self.overview["B8"] = self.status_counts[STATUS_PARTIAL]
        self.overview["B9"] = self.status_counts[STATUS_FAILED]
        self.overview["B10"] = round(self.total_duration / 3600, 3)
        self.overview["B11"] = round(
            self.total_duration / self.duration_count / 60, 3
        ) if self.duration_count else 0

    def maybe_save(self, force: bool = False) -> None:
        with self.lock:
            if not force and time.monotonic() - self.last_save < self.flush_seconds:
                return
            self._write_summary()
            temp = self.path.with_suffix(".tmp.xlsx")
            try:
                self.wb.save(temp)
                os.replace(temp, self.path)
                self.last_save = time.monotonic()
            except PermissionError:
                temp.unlink(missing_ok=True)
                logging.warning("Excel 台账正被其他程序占用；稍后重试保存。")

    def close(self) -> None:
        with self.lock:
            self.maybe_save(force=True)
            self.wb.close()


def task_status_for_year(year: int, actual_last: str) -> str:
    if year < utc_now().year:
        return STATUS_COMPLETED
    try:
        last_date = date.fromisoformat(actual_last[:10])
    except ValueError:
        return STATUS_PARTIAL
    return STATUS_COMPLETED if last_date >= date(year, 12, 31) else STATUS_PARTIAL


def download_task(
    task: Task,
    tracker: Tracker,
    lag_days: int,
    retries: int,
    parts_root: Path,
) -> tuple[str, str]:
    end_date = available_end_for_year(task.year, lag_days)
    if end_date is None:
        return task.task_id, "future"

    minimum_hours = expected_hours(task.year, end_date)
    try:
        _, first, last = validate_netcdf(task.output_file, minimum_hours)
        status = task_status_for_year(task.year, last)
        tracker.update(
            task.task_id,
            status=status,
            file_size_bytes=task.output_file.stat().st_size,
            coverage_start=first[:10],
            coverage_end=last[:10],
            output_file=str(task.output_file),
            last_error="",
        )
        if status == STATUS_COMPLETED:
            return task.task_id, "already_complete"
    except Exception:
        pass

    started = utc_now()
    started_clock = time.perf_counter()
    tracker.update(
        task.task_id,
        status=STATUS_DOING,
        start_time=iso_utc(started),
        end_time="",
        duration_seconds="",
        last_error="",
    )

    error: Exception | None = None
    for retry_index in range(retries):
        tracker.increment_attempts(task.task_id)
        try:
            client = cdsapi.Client()
            chunks = request_chunks(task.year, end_date, task.grid)
            task.output_file.parent.mkdir(parents=True, exist_ok=True)
            if len(chunks) == 1:
                retrieve_to_file(client, chunks[0][1], task.output_file)
            else:
                part_dir = parts_root / str(task.year) / task.grid.grid_id
                parts: list[Path] = []
                for chunk_name, request in chunks:
                    part = part_dir / f"{chunk_name}.nc"
                    if not part.exists():
                        retrieve_to_file(client, request, part)
                    else:
                        validate_netcdf(part)
                    parts.append(part)
                merge_netcdf(parts, task.output_file)

            _, first, last = validate_netcdf(task.output_file, minimum_hours)
            finished = utc_now()
            duration = round(time.perf_counter() - started_clock, 3)
            status = task_status_for_year(task.year, last)
            tracker.update(
                task.task_id,
                status=status,
                end_time=iso_utc(finished),
                duration_seconds=duration,
                file_size_bytes=task.output_file.stat().st_size,
                coverage_start=first[:10],
                coverage_end=last[:10],
                output_file=str(task.output_file),
                last_error="",
            )
            return task.task_id, status
        except Exception as exc:
            error = exc
            logging.exception("任务 %s 第 %d/%d 次尝试失败", task.task_id, retry_index + 1, retries)
            if retry_index + 1 < retries:
                time.sleep(min(60, 5 * (2 ** retry_index)))

    duration = round(time.perf_counter() - started_clock, 3)
    tracker.update(
        task.task_id,
        status=STATUS_FAILED,
        end_time=iso_utc(),
        duration_seconds=duration,
        last_error=str(error)[:2000] if error else "未知错误",
        force_save=True,
    )
    return task.task_id, STATUS_FAILED


def reconcile_existing(tasks: list[Task], tracker: Tracker, lag_days: int) -> None:
    logging.info("检查已有文件并恢复上次中断的状态……")
    for index, task in enumerate(tasks, start=1):
        old_status = str(tracker.value(task.task_id, "status") or STATUS_PENDING)
        end_date = available_end_for_year(task.year, lag_days)
        if task.output_file.exists() and end_date is not None:
            try:
                _, first, last = validate_netcdf(task.output_file)
                status = task_status_for_year(task.year, last)
                tracker.update(
                    task.task_id,
                    status=status,
                    file_size_bytes=task.output_file.stat().st_size,
                    coverage_start=first[:10],
                    coverage_end=last[:10],
                    output_file=str(task.output_file),
                    last_error="",
                )
            except Exception as exc:
                tracker.update(task.task_id, status=STATUS_PENDING, last_error=f"已有文件校验失败：{exc}")
        elif old_status in {STATUS_DOING, STATUS_COMPLETED, STATUS_PARTIAL}:
            tracker.update(task.task_id, status=STATUS_PENDING, last_error="上次状态未对应到有效文件")
        if index % 1000 == 0:
            logging.info("已检查 %d/%d 个任务", index, len(tasks))
    tracker.maybe_save(force=True)


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--start-year", type=int, default=DEFAULT_START_YEAR)
    parser.add_argument("--end-year", type=int, default=DEFAULT_END_YEAR)
    parser.add_argument("--output-dir", type=Path, default=script_dir / "era5_land_data")
    parser.add_argument("--tracker", type=Path, default=script_dir / "era5_land_download_tracker.xlsx")
    parser.add_argument("--workers", type=int, default=DEFAULT_WORKERS)
    parser.add_argument("--retries", type=int, default=DEFAULT_RETRIES)
    parser.add_argument("--availability-lag-days", type=int, default=DEFAULT_AVAILABILITY_LAG_DAYS)
    parser.add_argument("--flush-seconds", type=int, default=DEFAULT_FLUSH_SECONDS)
    parser.add_argument("--limit", type=int, default=None, help="仅运行前 N 个待处理任务，便于测试")
    parser.add_argument("--grid-id", action="append", help="仅下载指定格网，可重复，例如 --grid-id G0001")
    parser.add_argument("--init-only", action="store_true", help="只创建/更新台账，不开始下载")
    parser.add_argument("--skip-reconcile", action="store_true", help="跳过启动时的已有文件校验")
    parser.add_argument("--log-level", choices=["DEBUG", "INFO", "WARNING", "ERROR"], default="INFO")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    logging.basicConfig(
        level=getattr(logging, args.log_level),
        format="%(asctime)s %(levelname)s %(threadName)s %(message)s",
    )
    if args.start_year > args.end_year:
        raise SystemExit("start-year 不能大于 end-year")
    if args.workers < 1:
        raise SystemExit("workers 必须至少为 1")

    output_dir = args.output_dir.resolve()
    tracker_path = args.tracker.resolve()
    tasks = build_tasks(args.start_year, args.end_year, output_dir)
    if args.grid_id:
        requested = set(args.grid_id)
        tasks = [task for task in tasks if task.grid.grid_id in requested]
        unknown = requested - {task.grid.grid_id for task in tasks}
        if unknown:
            raise SystemExit(f"未知格网编号：{sorted(unknown)}")

    tracker = Tracker(tracker_path, tasks, args.availability_lag_days, args.flush_seconds)
    try:
        if not args.skip_reconcile:
            reconcile_existing(tasks, tracker, args.availability_lag_days)
        if args.init_only:
            logging.info("台账已初始化：%s", tracker_path)
            return 0

        candidates = []
        cutoff = (utc_now() - timedelta(days=args.availability_lag_days)).date()
        for task in tasks:
            status = str(tracker.value(task.task_id, "status") or STATUS_PENDING)
            if task.year > cutoff.year:
                continue
            if status in {STATUS_PENDING, STATUS_FAILED, STATUS_PARTIAL}:
                candidates.append(task)
        if args.limit is not None:
            candidates = candidates[: args.limit]
        logging.info("本次准备处理 %d 个任务，并发数 %d", len(candidates), args.workers)

        parts_root = output_dir / ".parts"
        with ThreadPoolExecutor(max_workers=args.workers, thread_name_prefix="era5") as pool:
            futures = {
                pool.submit(
                    download_task,
                    task,
                    tracker,
                    args.availability_lag_days,
                    args.retries,
                    parts_root,
                ): task
                for task in candidates
            }
            for completed_index, future in enumerate(as_completed(futures), start=1):
                task = futures[future]
                try:
                    task_id, status = future.result()
                    logging.info("[%d/%d] %s -> %s", completed_index, len(futures), task_id, status)
                except Exception:
                    logging.exception("任务 %s 出现未处理异常", task.task_id)
                tracker.maybe_save()
        return 0
    except KeyboardInterrupt:
        logging.warning("收到中断信号；正在保存 Excel 台账。")
        return 130
    finally:
        tracker.close()


if __name__ == "__main__":
    sys.exit(main())
